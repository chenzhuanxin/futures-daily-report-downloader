# -*- coding: utf-8 -*-
"""
期货公司日研报抓取 —— 通用工具模块
============================================================
提供 HTTP 请求、图片转 PDF、PDF 合并、docx 生成、路径工具等
被 fetchers.py 中 7 家公司的抓取函数复用。
"""
import datetime
import os
import re
import urllib.parse
import urllib3

import requests
from bs4 import BeautifulSoup
from PIL import Image
from pypdf import PdfReader, PdfWriter

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

UA = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
      '(KHTML, like Gecko) Chrome/120.0 Safari/537.36')

SESSION = requests.Session()
SESSION.headers.update({'User-Agent': UA})
VERIFY = False          # 部分期货公司站点证书链不完整，统一关闭证书校验
TIMEOUT = 60


def http_get(url, params=None, headers=None, timeout=TIMEOUT):
    """GET 请求，返回 requests.Response（自动处理编码）。"""
    h = dict(headers or {})
    r = SESSION.get(url, params=params, headers=h, timeout=timeout, verify=VERIFY)
    r.raise_for_status()
    return r


def http_download(url, dest_path, headers=None, timeout=180):
    """下载二进制文件到 dest_path。"""
    h = dict(headers or {})
    r = SESSION.get(url, headers=h, timeout=timeout, verify=VERIFY)
    r.raise_for_status()
    with open(dest_path, 'wb') as f:
        f.write(r.content)
    return r.content


def soup_of(response):
    """把响应文本解析为 BeautifulSoup 对象。"""
    return BeautifulSoup(response.text, 'lxml')


def images_to_pdf(image_paths, pdf_path):
    """将一张或多张图片合并转换为一个 PDF（每张图一页）。"""
    imgs = []
    for p in image_paths:
        im = Image.open(p)
        if im.mode != 'RGB':
            im = im.convert('RGB')
        imgs.append(im)
    first = imgs[0]
    rest = imgs[1:] if len(imgs) > 1 else []
    if rest:
        first.save(pdf_path, 'PDF', resolution=120.0, save_all=True, append_images=rest)
    else:
        first.save(pdf_path, 'PDF', resolution=120.0)
    for im in imgs:
        im.close()


def merge_pdfs(pdf_paths, out_path):
    """按顺序合并多个 PDF 为一个 PDF。"""
    writer = PdfWriter()
    for p in pdf_paths:
        reader = PdfReader(p)
        for page in reader.pages:
            writer.add_page(page)
    with open(out_path, 'wb') as f:
        writer.write(f)


def make_docx(title, body_paragraphs, out_path):
    """
    生成符合公文格式的 .docx：
      - 标题：黑体三号（16pt）加粗、居中
      - 正文：仿宋三号（16pt）、首行缩进 2 字符（32pt）、
        行间距固定值 30 磅、段前段后 0 磅
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    # 标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(title)
    run.bold = True
    run.font.name = '黑体'
    run.font.size = Pt(16)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    # 正文
    for text in body_paragraphs:
        para = doc.add_paragraph()
        pf = para.paragraph_format
        pf.first_line_indent = Pt(32)             # 首行缩进 2 字符（三号=16pt）
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(30)                  # 行距固定 30 磅
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        run = para.add_run(text)
        run.font.name = '仿宋'
        run.font.size = Pt(16)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

    doc.save(out_path)


def target_path(out_dir, company, date_str, ext):
    """返回目标文件完整路径：公司名+日期.后缀。"""
    filename = f'{company}{date_str}.{ext}'
    return os.path.join(out_dir, filename)


def ensure_dir(path):
    os.makedirs(path, exist_ok=True)
    return path


def today_str():
    return datetime.date.today().strftime('%Y-%m-%d')
